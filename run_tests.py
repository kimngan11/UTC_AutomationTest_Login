import os
import sys
import time
import argparse
import openpyxl
from selenium import webdriver
from selenium.webdriver.chrome.options import Options

# Đảm bảo mã hóa UTF-8 cho console Windows
sys.stdout.reconfigure(encoding='utf-8')

import config
from pages.login_page import LoginPage

# Danh sách ánh xạ 20 Test Cases theo file Excel TestCases_VanPhongDienTu_UTC.xlsx
TEST_CASES_MAPPING = [
    {
        "id": "TC01",
        "row": 5,
        "name": "Để trống Tên đăng nhập (chỉ nhập Mật khẩu)",
        "action": lambda page: test_tc01(page)
    },
    {
        "id": "TC02",
        "row": 6,
        "name": "Để trống Mật khẩu (chỉ nhập Tên đăng nhập)",
        "action": lambda page: test_tc02(page)
    },
    {
        "id": "TC03",
        "row": 7,
        "name": "Nhập đúng Tên đăng nhập, sai Mật khẩu",
        "action": lambda page: test_tc03(page)
    },
    {
        "id": "TC04",
        "row": 8,
        "name": "Nhập sai Tên đăng nhập, đúng định dạng Mật khẩu",
        "action": lambda page: test_tc04(page)
    },
    {
        "id": "TC05",
        "row": 9,
        "name": "Đăng nhập thành công và chọn 'Giữ tôi luôn đăng nhập'",
        "action": lambda page: test_tc05(page)
    },
    {
        "id": "TC06",
        "row": 10,
        "name": "Đăng nhập thành công và không chọn 'Giữ tôi luôn đăng nhập'",
        "action": lambda page: test_tc06(page)
    },
    {
        "id": "TC07",
        "row": 11,
        "name": "Để trống cả Tên đăng nhập và Mật khẩu",
        "action": lambda page: test_tc07(page)
    },
    {
        "id": "TC08",
        "row": 12,
        "name": "Kiểm tra tính năng ẩn mật khẩu (Masked characters)",
        "action": lambda page: test_tc08(page)
    },
    {
        "id": "TC09",
        "row": 13,
        "name": "Kiểm tra chức năng liên kết 'Đăng nhập bằng e-mail UTC'",
        "action": lambda page: test_tc09(page)
    },
    {
        "id": "TC10",
        "row": 14,
        "name": "Kiểm tra liên kết 'Bạn quên mật khẩu đăng nhập ?'",
        "action": lambda page: test_tc10(page)
    },
    {
        "id": "TC11",
        "row": 15,
        "name": "Đăng nhập bằng cách nhấn phím Enter từ bàn phím",
        "action": lambda page: test_tc11(page)
    },
    {
        "id": "TC12",
        "row": 16,
        "name": "Tên đăng nhập chỉ chứa toàn khoảng trắng (Space-only)",
        "action": lambda page: test_tc12(page)
    },
    {
        "id": "TC13",
        "row": 17,
        "name": "Mật khẩu chỉ chứa toàn khoảng trắng (Space-only)",
        "action": lambda page: test_tc13(page)
    },
    {
        "id": "TC14",
        "row": 18,
        "name": "Kiểm tra khả năng chống tấn công SQL Injection cơ bản",
        "action": lambda page: test_tc14(page)
    },
    {
        "id": "TC15",
        "row": 19,
        "name": "Kiểm tra khả năng chống tấn công XSS Script Injection",
        "action": lambda page: test_tc15(page)
    },
    {
        "id": "TC16",
        "row": 20,
        "name": "Nhập Tên đăng nhập có độ dài cực đại (300 ký tự)",
        "action": lambda page: test_tc16(page)
    },
    {
        "id": "TC17",
        "row": 21,
        "name": "Tên đăng nhập chứa tập ký tự đặc biệt (!@#$%^&*())",
        "action": lambda page: test_tc17(page)
    },
    {
        "id": "TC18",
        "row": 22,
        "name": "Kiểm tra văn bản gợi ý (Placeholder) trên các trường",
        "action": lambda page: test_tc18(page)
    },
    {
        "id": "TC19",
        "row": 23,
        "name": "Kiểm tra điều hướng chuyển con trỏ bằng phím Tab",
        "action": lambda page: test_tc19(page)
    },
    {
        "id": "TC20",
        "row": 24,
        "name": "Kiểm tra tính đúng đắn của các liên kết ở chân trang",
        "action": lambda page: test_tc20(page)
    }
]

# ----------------- Các hàm thực hiện kiểm thử -----------------

def test_tc01(page: LoginPage):
    page.enter_password("1256")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC01_empty_username.png")
    if "Bạn chưa nhập tên đăng nhập" in err:
        return True, f"Thông báo hiển thị đúng: '{err}'"
    return False, f"Thông báo không khớp: '{err}'"

def test_tc02(page: LoginPage):
    page.enter_username("huongnt")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC02_empty_password.png")
    if "Bạn chưa nhập mật khẩu" in err:
        return True, f"Thông báo hiển thị đúng: '{err}'"
    return False, f"Thông báo không khớp: '{err}'"

def test_tc03(page: LoginPage):
    page.enter_username("huongnt")
    page.enter_password("utc@235")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC03_wrong_password.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Thông báo lỗi hiển thị chính xác: '{err}'"
    return False, f"Thông báo lỗi không khớp: '{err}'"

def test_tc04(page: LoginPage):
    page.enter_username("huongthunguyen")
    page.enter_password("123456@utc")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC04_wrong_username.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Thông báo lỗi hiển thị chính xác: '{err}'"
    return False, f"Thông báo lỗi không khớp: '{err}'"

def test_tc05(page: LoginPage):
    page.enter_username(config.VALID_USERNAME)
    page.enter_password(config.VALID_PASSWORD)
    page.set_remember_me(True)
    is_checked = page.is_remember_me_checked()
    if not is_checked:
        return False, "Không thể tích chọn checkbox 'Giữ tôi luôn đăng nhập'"
    page.click_login()
    time.sleep(1.5)
    page.capture_screenshot("TC05_remember_me_checked.png")
    
    # Kiểm tra phản hồi thực tế từ hệ thống
    err = page.get_error_message(timeout=3)
    if err:
        return False, f"Thất bại: Hệ thống báo lỗi '{err}' (Tài khoản '{config.VALID_USERNAME}' không tồn tại hoặc sai mật khẩu trên máy chủ UTC)"
    return True, f"Đăng nhập thành công và ghi nhớ đăng nhập với tài khoản '{config.VALID_USERNAME}'"

def test_tc06(page: LoginPage):
    page.enter_username(config.VALID_USERNAME)
    page.enter_password(config.VALID_PASSWORD)
    page.set_remember_me(False)
    is_checked = page.is_remember_me_checked()
    if is_checked:
        return False, "Checkbox 'Giữ tôi luôn đăng nhập' đang bị chọn trái ý muốn"
    page.click_login()
    time.sleep(1.5)
    page.capture_screenshot("TC06_remember_me_unchecked.png")
    
    # Kiểm tra phản hồi thực tế từ hệ thống
    err = page.get_error_message(timeout=3)
    if err:
        return False, f"Thất bại: Hệ thống báo lỗi '{err}' (Tài khoản '{config.VALID_USERNAME}' không tồn tại hoặc sai mật khẩu trên máy chủ UTC)"
    return True, f"Đăng nhập thành công (không ghi nhớ) với tài khoản '{config.VALID_USERNAME}'"

def test_tc07(page: LoginPage):
    page.enter_username("")
    page.enter_password("")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC07_empty_all.png")
    if "Bạn chưa nhập tên đăng nhập" in err:
        return True, f"Thông báo hiển thị đúng: '{err}'"
    return False, f"Thông báo không khớp: '{err}'"

def test_tc08(page: LoginPage):
    is_masked = page.is_password_masked()
    page.capture_screenshot("TC08_password_masked.png")
    if is_masked:
        return True, "Trường mật khẩu đã được ẩn ký tự với type='password'"
    return False, "Trường mật khẩu không được ẩn ký tự (thiếu type='password')"

def test_tc09(page: LoginPage):
    href = page.get_google_login_href()
    page.capture_screenshot("TC09_google_sso.png")
    if "accounts.google.com" in href or "oauth2" in href:
        return True, f"Đường dẫn SSO Google UTC chính xác: {href[:65]}..."
    return False, f"Đường dẫn không hợp lệ: {href}"

def test_tc10(page: LoginPage):
    page.click_forgot_password()
    time.sleep(1.5)
    current_url = page.get_current_url()
    title = page.get_title()
    page.capture_screenshot("TC10_forgot_password.png")
    if "GetPass" in current_url and "Lấy lại mật khẩu" in title:
        return True, f"Đã chuyển hướng đến: {current_url} | Tiêu đề: '{title}'"
    return False, f"Chuyển hướng thất bại. URL: {current_url}, Title: '{title}'"

def test_tc11(page: LoginPage):
    page.press_enter_on_password("1256")
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC11_enter_key.png")
    if "Bạn chưa nhập tên đăng nhập" in err:
        return True, f"Kích hoạt gửi form bằng phím Enter thành công. Thông báo: '{err}'"
    return False, f"Thông báo không khớp: '{err}'"

def test_tc12(page: LoginPage):
    page.enter_username("   ")
    page.enter_password("123456")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC12_space_only_username.png")
    if "Tài khoản" in err or "mật khẩu" in err or "chưa nhập" in err:
        return True, f"Hệ thống từ chối username chỉ có khoảng trắng: '{err}'"
    return False, f"Xử lý không đúng: '{err}'"

def test_tc13(page: LoginPage):
    page.enter_username("huongnt")
    page.enter_password("   ")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC13_space_only_password.png")
    if "Tài khoản" in err or "mật khẩu" in err or "chưa nhập" in err:
        return True, f"Hệ thống từ chối password chỉ có khoảng trắng: '{err}'"
    return False, f"Xử lý không đúng: '{err}'"

def test_tc14(page: LoginPage):
    page.enter_username("' OR '1'='1")
    page.enter_password("123456")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC14_sql_injection.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Hệ thống an toàn với SQL Injection, từ chối đăng nhập: '{err}'"
    return False, f"Xử lý không an toàn hoặc lỗi: '{err}'"

def test_tc15(page: LoginPage):
    page.enter_username("<script>alert(1)</script>")
    page.enter_password("123456")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC15_xss_injection.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Hệ thống an toàn với XSS payload, từ chối đăng nhập: '{err}'"
    return False, f"Xử lý không an toàn hoặc lỗi: '{err}'"

def test_tc16(page: LoginPage):
    page.enter_username("a" * 300)
    page.enter_password("123456")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC16_long_username.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Hệ thống xử lý chuỗi 300 ký tự an toàn không bị tràn bộ đệm: '{err}'"
    return False, f"Xử lý không mong đợi: '{err}'"

def test_tc17(page: LoginPage):
    page.enter_username("admin!@#$%^&*()")
    page.enter_password("123456")
    page.click_login()
    time.sleep(1)
    err = page.get_error_message()
    page.capture_screenshot("TC17_special_chars.png")
    if "Tài khoản" in err or "mật khẩu" in err:
        return True, f"Hệ thống xử lý ký tự đặc biệt an toàn: '{err}'"
    return False, f"Xử lý không mong đợi: '{err}'"

def test_tc18(page: LoginPage):
    user_ph = page.get_username_placeholder()
    pwd_ph = page.get_password_placeholder()
    page.capture_screenshot("TC18_placeholders.png")
    if user_ph == "Tên đăng nhập" and pwd_ph == "Mật khẩu":
        return True, f"Placeholder hiển thị chuẩn: User='{user_ph}', Pass='{pwd_ph}'"
    return False, f"Placeholder không đúng: User='{user_ph}', Pass='{pwd_ph}'"

def test_tc19(page: LoginPage):
    focused = page.navigate_with_tab_key()
    page.capture_screenshot("TC19_tab_navigation.png")
    if focused == "userpwd":
        return True, "Phím Tab chuyển focus chính xác từ Tên đăng nhập sang Mật khẩu"
    return False, f"Focus không tới ô mật khẩu, đang ở: '{focused}'"

def test_tc20(page: LoginPage):
    help_href = page.get_help_center_href()
    fb_href = page.get_feedback_href()
    page.capture_screenshot("TC20_footer_links.png")
    if "hotrokythuat.utc.edu.vn" in help_href and "mailto:hotrokythuat@utc.edu.vn" in fb_href:
        return True, f"Liên kết chân trang chính xác: Help='{help_href}', Feedback='{fb_href}'"
    return False, f"Liên kết chân trang sai: Help='{help_href}', Feedback='{fb_href}'"

# ----------------- Cập nhật kết quả vào Excel -----------------

def update_excel_results(results):
    """Cập nhật kết quả vào file Excel TestCases_VanPhongDienTu_UTC.xlsx"""
    if not os.path.exists(config.EXCEL_FILE):
        print(f"[!] Không tìm thấy file Excel tại {config.EXCEL_FILE}")
        return

    try:
        wb = openpyxl.load_workbook(config.EXCEL_FILE)
        ws = wb["Test Cases"]
        for res in results:
            row_idx = res["row"]
            status_text = "Đạt (Pass)" if res["passed"] else "Thất bại (Fail)"
            ws.cell(row=row_idx, column=9, value=status_text)
        
        wb.save(config.EXCEL_FILE)
    except PermissionError:
        print(f"\n[!] CẢNH BÁO: File Excel '{os.path.basename(config.EXCEL_FILE)}' đang được mở trong phần mềm Excel!")
        print("    Vui lòng đóng cửa sổ Excel lại và chạy lại script để cập nhật kết quả tự động.")
    except Exception as e:
        print(f"\n[!] Lỗi khi ghi file Excel: {e}")

def create_driver(headless=False):
    options = Options()
    if headless:
        options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--ignore-certificate-errors")
    options.add_argument("--disable-notifications")
    return webdriver.Chrome(options=options)

def main():
    parser = argparse.ArgumentParser(description="Chạy kiểm thử tự động Văn phòng điện tử UTC")
    parser.add_argument("--headless", action="store_true", help="Chạy ở chế độ trình duyệt ẩn (headless)")
    parser.add_argument("--allure", action="store_true", help="Chạy kiểm thử với Pytest và tạo báo cáo Allure Report")
    args = parser.parse_args()

    is_headless = args.headless or config.HEADLESS

    if args.allure:
        print("[*] Chuyển tiếp sang bộ chạy Pytest + Allure Report...")
        import subprocess
        cmd = [sys.executable, "run_allure.py"]
        if is_headless:
            cmd.append("--headless")
        subprocess.run(cmd)
        return

    print("=" * 80)
    print(" BẮT ĐẦU CHẠY TOÀN BỘ BỘ KIỂM THỬ AUTOMATION: VĂN PHÒNG ĐIỆN TỬ UTC ")
    print(f" URL: {config.BASE_URL}")
    print(f" Chế độ trình duyệt: {'Headless (Ẩn)' if is_headless else 'UI (Trực quan)'}")
    print(f" Tổng số Test Cases: {len(TEST_CASES_MAPPING)}")
    print("=" * 80)

    driver = create_driver(headless=is_headless)
    page = LoginPage(driver)
    results = []

    passed_count = 0
    failed_count = 0

    try:
        for tc in TEST_CASES_MAPPING:
            tc_id = tc["id"]
            tc_name = tc["name"]
            print(f"\n>> Đang chạy [{tc_id}]: {tc_name}...")

            # Mở lại trang trước mỗi test case
            page.open()
            time.sleep(1)

            try:
                passed, detail = tc["action"](page)
            except Exception as e:
                passed = False
                detail = f"Ngoại lệ (Exception): {str(e)}"
                page.capture_screenshot(f"{tc_id}_error.png")

            status_str = "[PASS]" if passed else "[FAIL]"
            if passed:
                passed_count += 1
                print(f"   {status_str} -> {detail}")
            else:
                failed_count += 1
                print(f"   {status_str} -> {detail}")

            results.append({
                "id": tc_id,
                "name": tc_name,
                "row": tc["row"],
                "passed": passed,
                "detail": detail
            })

    finally:
        driver.quit()

    print("\n" + "=" * 80)
    print(" TỔNG KẾT KẾT QUẢ KIỂM THỬ ")
    print("=" * 80)
    print(f" Tổng số Test Cases: {len(results)}")
    print(f" Số ca ĐẠT (Pass)   : {passed_count}")
    print(f" Số ca HỎNG (Fail)  : {failed_count}")
    print(f" Tỷ lệ thành công   : {(passed_count / len(results) * 100):.1f}%")
    print(f" Ảnh chụp màn hình  : {config.SCREENSHOT_DIR}")
    print("=" * 80)

    # Ghi kết quả vào file Excel
    update_excel_results(results)

if __name__ == "__main__":
    main()
